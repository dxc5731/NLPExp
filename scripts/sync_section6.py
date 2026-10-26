# -*- coding: utf-8 -*-
"""修正 ⑥ 片段里第 3 条样例的描述，并把片段内容同步进完整报告 v0.2。"""
from copy import deepcopy
from docx import Document

ROOT = r"D:\agent-chat\dsh-deepseek\nlp\handoff"
FRAG = ROOT + r"\报告片段_⑥错误分析_ltm.docx"
FULL = ROOT + r"\nlp项目实验报告v0.2_ltm.docx"

OLD = "金标无实体，预测“发热”为 Symptom（误识别）"
NEW = "金标只标了“烧”，未标“发热”；预测多报了“发热”（误识别）"

# 1) 修正片段
frag = Document(FRAG)
fixed = 0
for t in frag.tables:
    for row in t.rows:
        for c in row.cells:
            if OLD in c.text:
                for p in c.paragraphs:
                    for r in p.runs:
                        if OLD in r.text:
                            r.text = r.text.replace(OLD, NEW)
                            fixed += 1
frag.save(FRAG)
print("片段修正处数:", fixed)

# 2) 用片段内容重建完整报告的 ⑥
frag = Document(FRAG)
body = []
started = False
for el in frag.element.body:
    tag = el.tag.split("}")[-1]
    if tag == "p":
        text = "".join(n.text or "" for n in el.iter() if n.tag.endswith("}t"))
        if text.strip().startswith("⑥") and not started:
            started = True
            continue
        if started:
            body.append(el)
    elif tag == "tbl" and started:
        body.append(el)

full = Document(FULL)
els = list(full.element.body)
i = next(k for k, e in enumerate(els)
         if e.tag.endswith("}p")
         and "".join(n.text or "" for n in e.iter() if n.tag.endswith("}t")).strip().startswith("⑥"))
j = next(k for k in range(i + 1, len(els))
         if els[k].tag.endswith("}p")
         and "".join(n.text or "" for n in els[k].iter() if n.tag.endswith("}t")).strip().startswith("⑦"))
anchor = els[i]
for e in els[i + 1:j]:
    e.getparent().remove(e)
for e in body:
    new = deepcopy(e)
    anchor.addnext(new)
    anchor = new
full.save(FULL)
print("已同步进完整报告")
