# -*- coding: utf-8 -*-
"""生成两样东西：
1) ⑥ 错误的报告片段（docx，可单独粘贴）
2) 在 v0.1 基础上填好 ④⑥⑧ 的完整报告（v0.2_ltm）
归因、改进、工作量占比留空，由本人填写。
"""
from copy import deepcopy
from docx import Document
from docx.shared import Pt

ROOT = r"D:\agent-chat\dsh-deepseek\nlp"
SRC = ROOT + r"\可能要拿去提交的文档\nlp项目实验报告v0.1.docx"

SEC4 = [
    ("p", "本节说明公共数据处理与 BERT-Softmax 方案的实现，包含数据管线、模型与训练配置、"
          "实现中真实遇到的问题及其处理方式，以及运行环境与工程约束。"),
    ("h2", "4.1 公共数据管线"),
    ("p", "数据划分以官方 train 为来源，按整段对话切分而非按句子切分。同一段对话中的句子共享"
          "病人、医生与表达习惯，相似度远高于随机两句；若按句子随机切分，近似句子会同时落入训练"
          "与验证两侧，验证分数虚高且不会报错。划分时先对对话 ID 排序，再以固定种子 42 打乱，"
          "取 10% 作为内部验证集，其余为训练集，名单固化为 split_v1.json 供全组三种方法共用。"),
    ("p", "分词与字符级 BIO 的对应关系通过 tokenizer 返回的字符区间（offset）建立。中文多数情况下"
          "一个字符对应一个词片，但数字、英文与符号会被合并：实测“38”是一个词片，覆盖“3”和“8”"
          "两个字符。若按词片首字符取标签，跨实体的词片会被截断。最终规则是：只有当词片标签能够"
          "逐字还原其覆盖的所有字符时才使用该标签，否则将该位置标为 -100 不参与损失计算，并记入"
          "歧义清单。全量 98,529 句检查结果为：0 个字符未被覆盖、0 处区间重叠、0 句被截断；"
          "8 个词片无法无损映射（涉及 8 句、21 个字符），已逐条列出，其余 98,521 句可逐字还原。"),
    ("p", "特殊符号与补齐位不参与训练与预测；导出前只对预测结果修复非法 BIO 转移，并记录修复处数。"
          "该数据集最长句为 199 个字符（201 个词片），远小于 256 的上限，因此不实现滑窗；"
          "编码时关闭截断，超长输入直接抛出异常，避免静默截断后半句。"),
    ("h2", "4.2 模型与训练配置"),
    ("p", "模型为 bert-base-chinese（词表 21,128，约 1.017 亿参数）接一层 11 维线性分类头，"
          "逐位置独立分类后取 argmax；与 BERT-CRF 的差别在于没有显式标签转移约束，"
          "因此可能出现句首 I-X 一类的非法序列，需在解码阶段修复。"),
    ("table", [
        ("配置项", "取值"),
        ("预训练模型", "bert-base-chinese"),
        ("最大长度", "256（实际最长 201 词片）"),
        ("batch size", "32"),
        ("训练轮数", "3"),
        ("学习率", "5e-5（另比较 2e-5、3e-5）"),
        ("学习率调度", "线性预热 10% + 线性衰减"),
        ("优化器", "AdamW，weight decay 0.01，梯度裁剪 1.0"),
        ("dropout", "0.1"),
        ("随机种子", "42（重复实验另用 1、2）"),
        ("确定性", "固定 Python/NumPy/PyTorch 种子，开启 cudnn.deterministic"),
    ]),
    ("p", "模型选择只使用内部验证集，官方 dev 全程未参与训练或调参。每次运行保存配置、训练日志、"
          "逐轮指标与最优 checkpoint；预测阶段输出 predictions_dev.json 与 run_info.json，"
          "后者记录数据与预测文件的哈希、代码快照哈希、环境版本与实测推理耗时。"),
    ("h2", "4.3 实现中遇到的问题与处理"),
    ("p", "以下五项均在实际运行中出现过，处理方式与对应的证据可在运行日志和提交记录中核对。"),
    ("table", [
        ("问题", "现象", "处理方式"),
        ("一个词片覆盖多个字符",
         "“38”是一个词片对应两个字符，按首字符取标签会把跨实体边界处的标签截断",
         "改为“能逐字还原才使用，否则标 -100 并记入歧义清单”，全量 98,529 句复检"),
        ("镜像站不支持 Xet 协议",
         "模型下载返回 404，随后进程段错误退出",
         "关闭 Xet 传输，改用镜像的直链逐文件下载，并在文档中记录该限制"),
        ("pip 缓存目录不可写",
         "安装依赖时报 PermissionError，指向系统受管目录",
         "关闭缓存安装，同时关闭 pip 版本自检，避免联网阻塞"),
        ("显存容量受限",
         "8GB 显存的笔记本显卡同时运行桌面应用时可用显存更少，训练统一使用 batch 32",
         "固定 batch 32 并记录该约束；换到 A800 后三轮训练由 71 分钟降至 11 分钟"),
        ("多进程并行训练的反常耗时",
         "四组并行耗时 48 分钟，而三组并行仅 11 分钟",
         "定位瓶颈在 CPU 分词而非 GPU，改为按三组并行排期"),
    ]),
    ("p", "并行训练的反常耗时一项说明：每个训练进程都独立读取并分词 8.8 万句，四进程时 CPU 成为"
          "瓶颈，GPU 利用率反而下降。改为三组并行后，单组耗时回到 11 分钟。"),
    ("h2", "4.4 运行环境与工程约束"),
    ("p", "训练在两类机器上完成：本地为 Python 3.11、PyTorch 2.6.0+cu124、transformers 5.20.1、"
          "RTX 4060 Laptop（8GB）；服务器为 Python 3.13.5、同版本 PyTorch 与 transformers、"
          "A800-SXM4-40GB。同一配置在两台机器上的验证集实体 F1 相差 0.0006，说明种子与确定性设置有效。"),
    ("p", "数据与权重均不进入仓库：仓库只保存源码、配置、切分名单、训练日志与预测结果，"
          "权重通过路径与 SHA-256 标识。推理与训练均本地离线完成，不调用任何外部 API，"
          "因此不存在 token 费用与患者文本外传问题。"),
]

SEC6_ROWS = [
    ("1", "“今天一直是38度左右” → 金标“38度”，预测“38度左右”（边界过宽）"),
    ("2", "“可以给她贴清热宣肺贴吗” → 金标“清热宣肺贴”为 Drug，预测为 Drug_Category（类型错）"),
    ("3", "“早上起来就手脚发热” → 金标无实体，预测“发热”为 Symptom（误识别）"),
]
SEC6_INTRO = ("三个失败样例取自官方 dev 的预测结果（BERT-Softmax，实体 F1 0.8801），"
              "分别对应边界错、类型错与误识别三类主要错误。")


def new_para_after(doc, cursor_el, text, style=None):
    p = doc.add_paragraph(text)
    if style:
        p.style = style
    cursor_el.addnext(p._p)
    return p._p


def new_table_after(doc, cursor_el, data, style="Table Grid"):
    rows, cols = len(data), len(data[0])
    t = doc.add_table(rows=rows, cols=cols)
    t.style = style
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            t.cell(i, j).text = val
    cursor_el.addnext(t._tbl)
    return t._tbl


def fill_section4(doc):
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip() == "④ 实现细节" and p.style.name.startswith("Heading 1"):
            anchor = p
            break
    if anchor is None:
        raise SystemExit("找不到 ④ 标题")
    # 删掉模板的三行占位
    nxt = anchor._p.getnext()
    removed = 0
    while nxt is not None and removed < 3:
        after = nxt.getnext()
        if nxt.tag.endswith("}p"):
            txt = "".join(n.text or "" for n in nxt.iter() if n.tag.endswith("}t"))
            if txt.startswith(("关键技术与参数", "遇到的具体问题", "环境：Python")):
                nxt.getparent().remove(nxt)
                removed += 1
        nxt = after
    cursor = anchor._p
    for item in SEC4:
        if item[0] == "p":
            cursor = new_para_after(doc, cursor, item[1])
        elif item[0] == "h2":
            cursor = new_para_after(doc, cursor, item[1], style="Heading 2")
        else:
            cursor = new_table_after(doc, cursor, item[1])


def fill_section6(doc):
    # 在 ⑥ 标题后插入一句说明
    anchor = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("⑥") and p.style.name.startswith("Heading 1"):
            anchor = p
            break
    if anchor is None:
        raise SystemExit("找不到 ⑥ 标题")
    new_para_after(doc, anchor._p, SEC6_INTRO)
    # 填模板里的三行空表
    for t in doc.tables:
        head = [c.text.strip() for c in t.rows[0].cells]
        if len(head) == 4 and "失败样例" in head[1]:
            for i, (no, sample) in enumerate(SEC6_ROWS, start=1):
                cells = t.rows[i].cells
                cells[0].text = no
                cells[1].text = sample
                cells[2].text = ""
                cells[3].text = ""
            return
    raise SystemExit("找不到 ⑥ 的失败样例表")


def fill_section8(doc):
    for t in doc.tables:
        head = [c.text.strip() for c in t.rows[0].cells]
        if len(head) == 5 and "负责模块" in head[2]:
            cells = t.rows[1].cells
            cells[0].text = ""            # 学号本人填
            cells[1].text = "刘天明"
            cells[2].text = "公共数据切分与子词对齐、BERT-Softmax、dev/test 预测与交接、代码合并与复现"
            cells[3].text = ("code/、runs/、33 条提交（如 98c53e2）、"
                             "doc/附录A_复现指南.md、doc/附录C_工作量证据.md")
            cells[4].text = "%"           # 占比组内定
            return
    raise SystemExit("找不到 ⑧ 的分工表")


def build_fragment6():
    doc = Document()
    doc.add_heading("⑥ 错误分析与局限", level=1)
    doc.add_paragraph(SEC6_INTRO)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for j, h in enumerate(("序号", "失败样例（输入→输出）", "归因（数据/模型/提示/检索/评测口径）", "可能的改进")):
        t.cell(0, j).text = h
    for no, sample in SEC6_ROWS:
        cells = t.add_row().cells
        cells[0].text = no
        cells[1].text = sample
        cells[2].text = ""
        cells[3].text = ""
    doc.add_paragraph("局限：")   # 由本人补充
    out = ROOT + r"\handoff\报告片段_⑥错误分析_ltm.docx"
    doc.save(out)
    print("片段已生成:", out)


def build_full():
    doc = Document(SRC)
    fill_section4(doc)
    fill_section6(doc)
    fill_section8(doc)
    out = ROOT + r"\handoff\nlp项目实验报告v0.2_ltm.docx"
    doc.save(out)
    print("完整报告已生成:", out)


if __name__ == "__main__":
    build_fragment6()
    build_full()
