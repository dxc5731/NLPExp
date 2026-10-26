# -*- coding: utf-8 -*-
"""读取一份 docx 的正文与表格，用于了解组里发来的模板要求。"""
import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)

print("=" * 60)
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"[{p.style.name}] {text}")

for ti, table in enumerate(doc.tables):
    print(f"\n--- 表格 {ti + 1}（{len(table.rows)} 行 × {len(table.columns)} 列）---")
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        print(" | ".join(cells))
